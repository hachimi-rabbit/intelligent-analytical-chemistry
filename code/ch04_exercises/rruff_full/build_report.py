"""Build the G04 full-RRUFF report deck from this module's outputs, on the G04 slide master.

Run from the repository root:
    python code/ch04_exercises/rruff_full/build_report.py
Writes resources/annual/2026-2027/student-reports/ch04/G04-Raman-Full-Report-new.pptx
"""
import json
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
REPORTS = ROOT / 'resources/annual/2026-2027/student-reports/ch04'
TEMPLATE = REPORTS / 'G04-Raman-Real-Report.pptx'
OUT = REPORTS / 'G04-Raman-Full-Report-new.pptx'
FONT = 'Noto Sans CJK SC'
FOOTER = 'G04 · 第4章机器学习 · 全量实测数据'
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY, BLUE, SLATE, MIST, PALE, OCHRE, SKY, GREY = (RGBColor.from_string(c) for c in
    ('15344C', '326A91', '556A79', 'CCD6DE', 'F1F5F7', 'A7651C', '8FB3CF', 'B8C2CA'))
MODELS = {'Dummy': '多数类', 'Cosine_1NN': '余弦近邻', 'PCA8_LR': 'PCA8＋逻辑回归', 'RBF_SVC': 'RBF-SVM'}
# Mineralogical relation of the most frequent cosine-1NN confusions shown on the error slide.
RELATIONS = {
    ('Amblygonite', 'Montebrasite'): ('锂磷铝石 → 羟磷锂铝石', 'F/OH 固溶体两端员'),
    ('Pyrope', 'Almandine'): ('镁铝榴石 → 铁铝榴石', '石榴子石族'),
    ('Grossular', 'Andradite'): ('钙铝榴石 → 钙铁榴石', '石榴子石族'),
    ('Tremolite', 'Actinolite'): ('透闪石 → 阳起石', '角闪石族，Mg/Fe 置换'),
    ('Richterite', 'Fluororichterite'): ('钠透闪石 → 氟钠透闪石', '角闪石族，OH/F 置换'),
    ('Fluoruvite', 'Dravite'): ('氟钙镁电气石 → 镁电气石', '电气石族'),
    ('Manganite', 'Groutite'): ('水锰矿 → 斜方水锰矿', 'MnOOH 同质多象'),
}


def load():
    metrics = json.loads((HERE / 'outputs/metrics.json').read_text())
    g = pd.read_csv(HERE / 'outputs/group_predictions.csv')
    folds = pd.read_csv(HERE / 'outputs/group_folds.csv')
    g['ok'] = g.true_mineral.eq(g.predicted_mineral)
    per_class = folds.groupby('mineral').validation_group.nunique()
    single = set(per_class[per_class.eq(1)].index)
    g['support'] = g.true_mineral.map(per_class).clip(upper=6)
    return metrics, g, per_class, single


metrics, groups, per_class, single = load()
M = metrics['metrics']
fold_acc = groups.groupby(['model', 'fold']).ok.mean().unstack()
by_support = groups.groupby(['model', 'support']).ok.mean().unstack()
support_n = groups[groups.model.eq('Cosine_1NN')].support.value_counts().sort_index()
cos = groups[groups.model.eq('Cosine_1NN')]
cos_err = cos[~cos.ok]
to_single = int(cos_err.predicted_mineral.isin(single).sum())
pairs = cos_err.groupby(['true_mineral', 'predicted_mineral']).size()
svm = groups[groups.model.eq('RBF_SVC')].set_index('validation_group')
agree = float(cos.set_index('validation_group').predicted_mineral.eq(svm.predicted_mineral).mean())
hist = per_class.clip(upper=5).value_counts().sort_index()


def demo_facts():
    """What the two demo pages display for the clicks scripted on the demo slide; fails if the pages' data changed."""
    page = json.loads((HERE / 'outputs/browser-data.js').read_text().split('=', 1)[1].rstrip().rstrip(';'))
    names = page['classes']
    quartz = [r for r in page['spectra'] if r['mineral'] == 'Quartz']
    first = [r for r in quartz if r['specimen'] == quartz[0]['specimen']]
    cosine = [names[r['predictions']['Cosine_1NN']] for r in first]
    dummy = {names[r['predictions']['Dummy']] for r in quartz}
    abelsonite = [r for r in page['spectra'] if r['mineral'] == 'Abelsonite']
    sim = json.loads((ROOT / 'assets/js/ch04-results.js').read_text().split('var data=', 1)[1].split('; if(typeof', 1)[0])
    shared = {split: sum(len({x == 1 for x in r['folds']}) > 1 for r in sim['folds'][split]) for split in ('row_random', 'specimen_group')}
    assert names[0] == 'Abelsonite' and len(abelsonite) == 1 and abelsonite[0]['fold'] < 0
    assert [r['laser'] for r in first] == [514, 532, 780] and cosine == ['Sodalite', 'Quartz', 'Quartz'] and first[0]['fold'] == 0
    assert dummy == {'Grossular'} and shared == {'row_random': 44, 'specimen_group': 0}
    return first[0]['specimen']


QUARTZ_SPECIMEN = demo_facts()


def acc(model):
    return M[model]['group']['accuracy']


def pct(x):
    return f'{100 * x:.1f}%'


prs = Presentation(TEMPLATE)
ids = prs.slides._sldIdLst
for sld in list(ids):
    prs.part.drop_rel(sld.rId)
    ids.remove(sld)
BLANK = next(l for l in prs.slide_layouts if l.name == 'Blank')


def set_typeface(rpr):
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set('typeface', FONT)


def font(run, size, color=NAVY, bold=False):
    run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, color
    set_typeface(run._r.get_or_add_rPr())


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, space=6):
    """paras: (text, size, color, bold) per paragraph, or a list of such runs for one paragraph."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment, p.space_after = align, Pt(space)
        for t, size, color, bold in (para if isinstance(para, list) else [para]):
            r = p.add_run()
            r.text = t
            font(r, size, color, bold)
    return box


def rect(slide, x, y, w, h, color, rounded=True):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def new_slide(title, notes):
    slide = prs.slides.add_slide(BLANK)
    text(slide, 0.67, 0.42, 12.0, 0.85, [(title, 33, NAVY, True)], anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 0.67, 7.02, 10.94, 0.26, [(FOOTER, 12, SLATE, False)])
    text(slide, 12.19, 7.02, 0.52, 0.26, [(str(len(prs.slides)), 12, SLATE, False)], align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def stat(slide, x, y, w, value, label, color, size=48):
    text(slide, x, y, w, 0.95, [(value, size, color, True)])
    text(slide, x, y + 1.0, w, 0.8, [(label, 16.5, SLATE, False)])


def chart(slide, kind, x, y, w, h, categories, series, colors, title, fmt='0.000', vmin=None, vmax=None, legend=False, positions=None):
    """positions: for line charts, one list of XL_LABEL_POSITION per series (one entry per point)."""
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    ch = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    ch.font.size, ch.font.color.rgb = Pt(14), SLATE
    set_typeface(ch.font._rPr)
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    font(ch.chart_title.text_frame.paragraphs[0].runs[0], 16, NAVY, True)
    ch.has_legend = legend
    if legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        ch.legend.font.size = Pt(14)
        set_typeface(ch.legend.font._rPr)
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = MIST
    va.format.line.fill.background()
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = {'0': '0', '0%': '0%'}.get(fmt, '0.0'), False
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    ch.category_axis.format.line.color.rgb = MIST
    plot = ch.plots[0]
    if kind != XL_CHART_TYPE.LINE_MARKERS:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format, dl.number_format_is_linked = fmt, False
        dl.font.size, dl.font.color.rgb = Pt(14), NAVY
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        plot.gap_width = 70
    for i, (s, c) in enumerate(zip(plot.series, colors)):
        if kind == XL_CHART_TYPE.LINE_MARKERS:
            s.format.line.color.rgb, s.format.line.width = c, Pt(3.5)
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = c
            s.marker.format.line.color.rgb = c
            s.marker.size = 9
            s.smooth = False
            pos = positions[i] if positions else [XL_LABEL_POSITION.ABOVE] * len(categories)
            for j, point in enumerate(s.points):
                label = point.data_label
                label.position = pos[j]
                label.font.size, label.font.color.rgb = Pt(14), c
                number_format(label._dLbl, fmt)
        else:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = c
    return ch


def number_format(dLbl, fmt):
    nf = dLbl.makeelement(qn('c:numFmt'), {'formatCode': fmt, 'sourceLinked': '0'})
    dLbl.find(qn('c:idx')).addnext(nf)


def table(slide, x, y, col_w, rows, row_h=0.52, size=16):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                   Inches(sum(col_w)), Inches(row_h * len(rows)))
    tb = shape.table
    for j, cw in enumerate(col_w):
        tb.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tb.rows[i].height = Inches(row_h)
        for j, value in enumerate(row):
            cell = tb.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (PALE if i % 2 else WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.1)
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = value
            font(r, size, WHITE if i == 0 else NAVY, i == 0)
    return tb


# 1 Title
s = new_slide('全量实测拉曼光谱的矿物分类',
    '【本讲稿逐字照念即可；方括号里是动作提示，不要念出来。全程约 10 分钟。】\n'
    '大家好，我们是 G04 组。今天汇报的题目是：全量实测拉曼光谱的矿物分类。\n'
    '上一次汇报，我们只用了三种矿物：石英、方解石和石膏，一共 52 条谱、22 个标本。'
    '三种方法都做到了百分之百，所以分不出方法的好坏。\n'
    f'这一次，我们把公开数据库里的全部实测谱都用上了：{metrics["source_spectra"]:,} 条谱，'
    f'{metrics["source_specimens"]:,} 个标本，{metrics["source_classes"]:,} 种矿物。\n'
    '我们要回答的问题和上次一样：模型能不能认出一个它从来没见过的标本。\n'
    '【翻页】')
text(s, 0.67, 2.0, 11.98, 1.6, [('从 3 类到 1,669 类，模型还能认出新标本吗？', 38.25, NAVY, True)], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.71, 4.2, 11.46, 2.2, [
    (f'{metrics["source_spectra"]:,} 条实测谱 · {metrics["source_specimens"]:,} 个标本 · {metrics["source_classes"]:,} 类矿物', 24.75, BLUE, False),
    ('同一套按标本留出的验证规则，四种固定参数的方法', 22.5, NAVY, False),
    ('谱图取自 RRUFF 数据库（DTU 存档 v1），已经来源方预处理', 19.5, SLATE, False)], space=14)

# 2 Answer first
s = new_slide('类别扩大到 1,669 类后，正确率降至约七成',
    '先说结论。\n'
    f'【指向左边两个大数字】只有三种矿物的时候，正确率是百分之百；扩大到 {metrics["usable_classes"]:,} 种矿物以后，'
    f'最好的方法 SVM 降到了 {pct(acc("RBF_SVC"))}。\n'
    f'【指向右边方框】我们有三个发现。第一，最简单的余弦近邻得到 {pct(acc("Cosine_1NN"))}，和 SVM 基本持平。'
    '第二，错误主要发生在成分相近的同族矿物之间。第三，一种矿物的独立标本越多，模型就认得越准。\n'
    '【指向左下方两行字】先说明一下怎么打分。同一个标本的全部光谱一起留出来考试，投票得到一个答案，这叫一个验证组。'
    f'至少有两个独立标本的 {metrics["evaluated_classes"]} 种矿物才能被验证，'
    f'其余 {metrics["training_only_classes"]:,} 种只放在参考库里。\n'
    '后面几页依次展开。【翻页】')
stat(s, 0.67, 1.7, 3.2, '100%', '三类矿物，22 个标本', SLATE, 54)
text(s, 3.55, 1.82, 0.8, 0.9, [('→', 44, GREY, False)])
stat(s, 4.4, 1.7, 3.6, pct(acc('RBF_SVC')), '1,669 类 · SVM', BLUE, 54)
rect(s, 8.55, 1.55, 4.12, 4.95, PALE)
text(s, 8.85, 1.8, 3.6, 4.5, [
    ('三个发现', 23.25, OCHRE, True),
    (f'余弦近邻 {pct(acc("Cosine_1NN"))}，与 SVM 基本持平', 18, NAVY, False),
    ('错误多在同族、固溶体矿物之间', 18, NAVY, False),
    ('独立标本越多，正确率越高', 18, NAVY, False)], space=16)
text(s, 0.67, 4.35, 7.5, 2.2, [
    ('评分单位是“验证组”：同一标本的全部谱一起留出、投票得到一个预测。', 18, NAVY, False),
    ('只有至少两个独立标本的 606 类能被验证；其余 1,063 类只作参考。', 18, SLATE, False)], space=12)

# 3 Data
s = new_slide('数据多，但多数矿物只有一个标本',
    f'这一页是数据。【指向左上方数字】{metrics["source_spectra"]:,} 条记录我们逐条检查过，'
    f'其中 {metrics["excluded_spectra"]} 条的标签互相冲突，比如同一个标本有两个矿物名，我们把它们隔离了。'
    f'剩下 {metrics["usable_spectra"]:,} 条可用，覆盖 {metrics["usable_classes"]:,} 种矿物。\n'
    f'【指向右边图里灰色的柱子】关键是这根灰色柱子：{hist.iloc[0]:,} 种矿物只有一个标本。'
    '只有一个标本，就没法检验模型能不能认出同种矿物的新标本，所以它们只能放进参考库，不产生分数。\n'
    f'【指向左下方橙色数字】真正能验证的是 {metrics["evaluated_classes"]} 种矿物。\n'
    '要注意，参考库里的矿物虽然不打分，但模型仍然可能把别的谱错判成它们，这种错误照样算错。【翻页】')
stat(s, 0.67, 1.6, 3.0, f'{metrics["usable_spectra"]:,}', f'条谱可用（隔离 {metrics["excluded_spectra"]} 条标签冲突）', NAVY, 44)
stat(s, 0.67, 3.25, 3.0, f'{metrics["usable_classes"]:,}', '类矿物进入模型', NAVY, 44)
stat(s, 0.67, 4.9, 3.0, f'{metrics["evaluated_classes"]}', '类有 ≥2 个标本，可以验证', OCHRE, 44)
ch = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 3.9, 1.45, 8.77, 5.0, ['1 个', '2 个', '3 个', '4 个', '≥5 个'],
           [('矿物种数', [int(v) for v in hist.values])], [BLUE], '每种矿物拥有的独立标本数', fmt='0')
pt = ch.plots[0].series[0].points[0]
pt.format.fill.solid()
pt.format.fill.fore_color.rgb = GREY
text(s, 3.9, 6.5, 8.77, 0.4, [(f'灰色 {hist.iloc[0]:,} 类只有一个标本：只进参考库，不产生验证分数', 15, SLATE, False)])

# 4 Validation
s = new_slide('五折分组验证：同一标本不会同时出现在训练和验证中',
    f'这一页讲怎么验证。【指向五个色块】我们把能验证的 {metrics["evaluated_groups"]:,} 个验证组分成五个分区。'
    '每一轮拿一个分区来考试，其余四个用来训练，轮流做五次。'
    '同一个标本的全部光谱永远在同一边，不会一半拿去训练、一半拿去考试。\n'
    f'【指向灰色长条】{metrics["training_only_classes"]:,} 种只有一个标本的矿物，每一轮都只放在训练这边，当作参考。\n'
    '【指向下方表格】我们比较了四种方法。余弦近邻，就是找形状最像的一条训练谱，沿用它的矿物名。'
    'PCA 八维加逻辑回归，先把每条光谱压缩成 8 个数，再分类。RBF 核的 SVM，学习非线性的分类边界。'
    '还有一个多数类基线，它不看光谱，永远猜最常见的那种矿物。\n'
    '四种方法的参数都是事先定好的，没有根据考试成绩去调。【翻页】')
text(s, 0.71, 1.55, 11.9, 0.5, [('每轮留出一个分区；同一标本的全部谱始终在同一侧', 21, NAVY, False)])
for i in range(5):
    x = 0.71 + i * 2.42
    rect(s, x, 2.3, 2.2, 1.0, OCHRE if i == 0 else BLUE)
    text(s, x, 2.3, 2.2, 1.0, [(f'分区 {i + 1}', 21, WHITE, True), ('本轮验证' if i == 0 else '本轮训练', 16.5, WHITE, False)],
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, space=2)
rect(s, 0.71, 3.55, 11.88, 0.75, PALE)
text(s, 1.0, 3.55, 11.4, 0.75, [(f'参考库：{metrics["training_only_classes"]:,} 类单标本矿物，每一轮都只参与训练', 18, SLATE, False)],
     anchor=MSO_ANCHOR.MIDDLE)
table(s, 0.71, 4.65, [3.4, 8.48], [
    ['方法（参数固定）', '怎样判别'],
    ['余弦近邻', '找形状最像的一条训练谱，沿用它的矿物名'],
    ['PCA8＋逻辑回归', '先压缩到 8 个主成分，再学类别边界'],
    ['RBF-SVM', '学习类别之间的非线性边界；多数类基线不看光谱']], row_h=0.48, size=16)

# 5 Results
s = new_slide('余弦近邻与 SVM 打平，压到 8 维的线性模型明显落后',
    '这是四种方法的结果。【指向图】蓝色柱子是正确率，橙色柱子是宏 F1，也就是每种矿物权重相同的平均成绩。\n'
    f'【指向最左边一组柱子】多数类基线只有 {pct(acc("Dummy"))}，说明这个任务没法靠猜。\n'
    f'【指向第三组柱子】PCA 加逻辑回归是 {pct(acc("PCA8_LR"))}，明显落后。原因是只留了 8 维，'
    '面对上千种矿物，信息丢得太多。这并不说明线性模型本身不行。\n'
    f'【指向第二组和第四组柱子】余弦近邻是 {pct(acc("Cosine_1NN"))}，SVM 是 {pct(acc("RBF_SVC"))}，只差大约一个百分点。\n'
    f'【指向右边文字】而且在五个分区上，余弦近邻在 {fold_acc.loc["Cosine_1NN"].min():.1%} 到 {fold_acc.loc["Cosine_1NN"].max():.1%} 之间，'
    f'SVM 在 {fold_acc.loc["RBF_SVC"].min():.1%} 到 {fold_acc.loc["RBF_SVC"].max():.1%} 之间。'
    '分区之间的波动比两种方法的差距还大，所以我们只能说两者相当，不能说 SVM 更好。'
    f'另外，两种方法在 {agree:.0%} 的验证组上给出了相同的答案。【翻页】')
chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 0.67, 1.45, 8.3, 5.3, list(MODELS.values()),
      [('组级正确率', [acc(m) for m in MODELS]), ('宏 F1', [M[m]['group']['macro_f1'] for m in MODELS])],
      [BLUE, OCHRE], f'{metrics["evaluated_classes"]} 类、{metrics["evaluated_groups"]:,} 组的折外结果', fmt='0.000', vmin=0, vmax=1, legend=True)
text(s, 9.3, 1.7, 3.37, 4.9, [
    ('差距小于波动', 23.25, BLUE, True),
    (f'五个分区上，近邻 {fold_acc.loc["Cosine_1NN"].min():.1%}–{fold_acc.loc["Cosine_1NN"].max():.1%}，'
     f'SVM {fold_acc.loc["RBF_SVC"].min():.1%}–{fold_acc.loc["RBF_SVC"].max():.1%}。', 18, NAVY, False),
    ('两者多数时候一致', 23.25, BLUE, True),
    (f'{agree:.0%} 的验证组给出相同答案。', 18, NAVY, False)], space=12)

# 6 Support
cos_s, svm_s = by_support.loc['Cosine_1NN'], by_support.loc['RBF_SVC']
s = new_slide('同一种矿物的标本越多，认得越准',
    '这一页看标本数量的影响。【指向横轴】横轴是这种矿物一共有几个独立标本，纵轴是正确率。\n'
    f'【指向两条线的最左端】只有 2 个标本的时候，训练这边只剩 1 个可以参考，'
    f'余弦近邻的正确率是 {cos_s.iloc[0]:.0%}，SVM 是 {svm_s.iloc[0]:.0%}。\n'
    f'【指向两条线的最右端】有 6 个以上标本的时候，余弦近邻升到 {cos_s.iloc[-1]:.0%}，SVM 升到 {svm_s.iloc[-1]:.0%}。\n'
    '两条线都在往上走。这说明瓶颈主要在数据：为同一种矿物多收集来自不同产地的标本，比换一个更复杂的模型更有用。【翻页】')
cats = [f'{k} 个' if k < 6 else '≥6 个' for k in support_n.index]
# Label the upper curve above and the lower curve below at each point so labels do not collide.
upper = by_support.loc['Cosine_1NN'].to_numpy() > by_support.loc['RBF_SVC'].to_numpy()
ABOVE, BELOW = XL_LABEL_POSITION.ABOVE, XL_LABEL_POSITION.BELOW
chart(s, XL_CHART_TYPE.LINE_MARKERS, 0.67, 1.45, 8.3, 5.3, cats,
      [('余弦近邻', list(by_support.loc['Cosine_1NN'])), ('RBF-SVM', list(by_support.loc['RBF_SVC']))],
      [SKY, BLUE], '按该矿物的独立标本数分组的正确率', fmt='0%', vmin=0.5, vmax=0.9, legend=True,
      positions=[[ABOVE if u else BELOW for u in upper], [BELOW if u else ABOVE for u in upper]])
lo, hi = by_support.loc['Cosine_1NN'].iloc[0], by_support.loc['Cosine_1NN'].iloc[-1]
stat(s, 9.3, 1.7, 3.37, f'{lo:.0%} → {hi:.0%}', '余弦近邻：2 个标本 → 6 个以上', BLUE, 40)
text(s, 9.3, 3.8, 3.37, 2.8, [
    ('各点的验证组数', 18, SLATE, True),
    ('，'.join(f'{c} {n}' for c, n in zip(cats, support_n.values)), 16.5, SLATE, False)], space=8)

# 7 Errors
rows = [['真实矿物 → 被判为', '关系', '组数']]
for (t, p), n in pairs.sort_values(ascending=False).items():
    if (t, p) in RELATIONS:
        rows.append([RELATIONS[t, p][0], RELATIONS[t, p][1], str(int(n))])
    if len(rows) == 7:
        break
assert rows[1][0] == RELATIONS['Amblygonite', 'Montebrasite'][0]
s = new_slide('错误集中在成分相近的同族矿物之间',
    '那模型错在哪里？【指向表格】这是余弦近邻最常见的几种错误。'
    '第一行，锂磷铝石被判成了羟磷锂铝石。它们是同一个固溶体系列的两端，只差氟和羟基。\n'
    '下面几行也一样，都是石榴子石族、角闪石族、电气石族里的同族矿物。它们结构相同，只是部分离子互相替换，'
    '拉曼峰的位置只差几个波数。这类错误在矿物学上是合理的，只靠这一段光谱很难分开，需要再加上化学成分等信息。\n'
    f'【指向右边橙色数字】另外，余弦近邻一共错了 {len(cos_err)} 组，其中 {to_single} 组，也就是 {to_single / len(cos_err):.0%}，'
    '被判成了只有一个标本的矿物。参考库里的这些矿物没法验证，却会参与竞争，抢走答案。【翻页】')
table(s, 0.67, 1.6, [4.3, 3.3, 0.9], rows, row_h=0.6, size=16)
text(s, 0.67, 5.95, 8.5, 0.5, [('余弦近邻最常见的混淆对（按组数）', 15, SLATE, False)])
rect(s, 9.4, 1.6, 3.27, 4.2, PALE)
stat(s, 9.7, 1.9, 2.8, f'{to_single / len(cos_err):.0%}', f'的错误（{to_single}/{len(cos_err)} 组）被判成只有一个标本的矿物', OCHRE, 48)
text(s, 9.7, 4.3, 2.8, 1.4, [('参考库里的“孤本”无法验证，却会参与竞争', 16.5, NAVY, False)])

# 8 Limits and directions
s = new_slide('适用边界与改进方向',
    '这一页讲结果的适用边界和改进方向。\n'
    f'【指向左栏】边界有三条。第一，{metrics["training_only_classes"]:,} 种矿物只有一个标本，我们的结果不代表这些矿物。'
    '第二，为了让所有光谱有共同的输入，只用了波数 304 到 852 这一段，更低和更高波数的峰没有用上。'
    '第三，这是数据库内部的验证，没有用其他实验室、其他仪器测的光谱检验过。\n'
    '【指向右栏】对应的改进方向也有三条。第一，为同一种矿物收集更多不同产地的标本。'
    '第二，按照 2024 年发表的 REFORMS 规范，完整报告数据划分方式和复现信息。'
    '第三，2025 年 Lilek 等人的研究表明，验证方式会改变模型的排名；下一步也可以试 TabPFN 这类新的基线方法，这个我们这次还没有做。\n'
    '【指向页面底部】参考文献列在页面下方。【翻页】')
text(s, 0.67, 1.6, 5.8, 5.0, [
    ('边界', 23.25, OCHRE, True),
    (f'{metrics["training_only_classes"]:,} 类只有一个标本，没有验证分数', 18, NAVY, False),
    ('只用共同区间 304–852 cm⁻¹，其余波段的峰未参与', 18, NAVY, False),
    ('数据库内验证，未做外部实验室或跨仪器测试', 18, NAVY, False)], space=16)
text(s, 6.87, 1.6, 5.8, 5.0, [
    ('方向', 23.25, BLUE, True),
    ('为同一矿物收集更多不同产地的独立标本', 18, NAVY, False),
    ('按 REFORMS（2024）报告划分方式与复现信息', 18, NAVY, False),
    ('验证方式会改变模型排序（Lilek 等，2025）；可尝试 TabPFN（2025）等新基线', 18, NAVY, False)], space=16)
text(s, 0.67, 5.75, 12.0, 1.0, [
    ('Kapoor 等，Sci. Adv. 2024，doi:10.1126/sciadv.adk3452 · Lilek 等，J. Raman Spectrosc. 2025，doi:10.1002/jrs.6842 · '
     'Hollmann 等，Nature 2025，doi:10.1038/s41586-024-08328-6', 12, SLATE, False),
    ('数据：Li、Schmidt、Alstrøm，DTU 数据存档 v1，doi:10.11583/DTU.20222331.v1，CC BY 4.0；原始数据库 RRUFF', 12, SLATE, False)], space=4)

# 9 Demo
s = new_slide('现场演示：折外预测、参考库与数据泄漏',
    '【提前在浏览器里打开两个页面：assets/interactive/ch04-full.html 和 assets/interactive/ch04.html。】\n'
    '下面做一个两分钟的现场演示。【切换到 ch04-full.html 页面，不要动任何下拉框】\n'
    f'第一步。页面默认选中的是石英，模型是余弦近邻。【指向“预测”那一行字】这是石英标本 {QUARTZ_SPECIMEN} 用 514 纳米激光测的谱。'
    '页面显示：余弦近邻预测 Sodalite，也就是方钠石，与参考标签不一致。单看这一条谱，模型判错了。\n'
    '【在“实测光谱”下拉框里选第二项，532 nm】换成同一个标本用 532 纳米测的谱，预测是 Quartz，判对了。'
    '【再选第三项，780 nm】780 纳米的谱也判对了。\n'
    '这个标本三条谱，两条判对，投票结果就是石英。这就是我们按标本投票打分的原因。\n'
    '【在“固定模型”下拉框里选“多数类基线”】换成多数类基线，它预测 Grossular，也就是钙铝榴石。'
    '不管什么谱它都猜这一种，因为这是数据里最常见的矿物。\n'
    '【在“矿物类别”下拉框里选第一项 Abelsonite】第二步。这是卟啉镍石，它只有一个标本。'
    '页面显示：本类别只有一个验证组，本谱只参与训练，没有验证预测。这就是前面说的参考库。\n'
    '【切换到 ch04.html 页面】第三步。这个页面用的是模拟数据，用来说明为什么要按标本划分。\n'
    '【“划分方式”选“逐谱随机划分（反例）”，“验证折”选 1】如果把光谱打乱随机划分，第 1 折里，训练和验证两边共享了 44 个标本。'
    '同一个标本既在训练里，又在考试里，这就是数据泄漏。\n'
    '【“划分方式”改选“按物理标本分组”】按标本分组以后，共享的标本数变成 0。\n'
    '注意，这一页是模拟数据，它的结果不算进我们的实测成绩。\n'
    '【如果网页打不开：不做操作，指着本页三个方框，把每个方框里的两行字依次念出来】\n'
    '【切回 PPT，翻页】')
for i, (head, body) in enumerate([
        ('① 全量页：石英（Quartz）', '默认的 514 nm 谱被判成方钠石；同一标本另两条谱判对，投票结果为石英'),
        ('② 选择 Abelsonite（卟啉镍石）', '只有一个标本：只在参考库中，没有验证预测'),
        ('③ 模拟页：切换两种划分', '逐谱随机划分时，第 1 折两侧共享 44 个标本；按标本分组后为 0')]):
    y = 1.65 + i * 1.6
    rect(s, 0.71, y, 11.9, 1.3, PALE)
    text(s, 1.0, y + 0.18, 11.3, 1.0, [(head, 23.25, BLUE, True), (body, 18, NAVY, False)], space=6)

# 10 Summary
spread = [fold_acc.loc[m].max() - fold_acc.loc[m].min() for m in ('Cosine_1NN', 'RBF_SVC')]
s = new_slide('总结',
    '最后总结三点。\n'
    f'【指向第一条】第一，类别从 3 种扩大到 {metrics["usable_classes"]:,} 种以后，同样按标本留出验证，最好的正确率大约是 71%。\n'
    '【指向第二条】第二，简单方法就够用。余弦近邻和 SVM 相当，而且错误大多发生在同族矿物之间。\n'
    '【指向第三条】第三，数据是瓶颈。每种矿物的独立标本越多，认得越准，这比换模型更有效。\n'
    '【指向底部方框】下一步可以做两件事：为同一种矿物收集更多独立标本，以及用其他实验室、其他仪器测的光谱做外部检验。\n'
    '我们的汇报就到这里，谢谢大家。\n'
    '\n'
    '【以下是可能被问到的问题和回答，被问到时照念对应的回答。】\n'
    '问：SVM 的正确率更高，为什么说两者相当？\n'
    '答：两者只差大约一个百分点；而同一种方法在不同分区之间，正确率就能差 '
    f'{100 * min(spread):.0f} 到 {100 * max(spread):.0f} 个百分点。'
    '差距在波动范围之内，我们也没有做显著性检验，所以不能说谁更好。\n'
    '问：为什么只用 304 到 852 这一段波数？\n'
    '答：不同光谱的测量范围不一样，304 到 852 是所有光谱都覆盖的共同区间。我们没有外推或者补造缺失的数据。\n'
    '问：只有一个标本的矿物为什么不打分？\n'
    '答：我们要检验的是模型能不能认出没见过的标本。只有一个标本的矿物，把它拿去考试以后，训练里就没有同种矿物了，没法检验。\n'
    '问：为什么不调参数？\n'
    '答：参数事先固定，是为了避免根据考试成绩挑参数，让成绩虚高。\n'
    '问：为什么不用更复杂的模型？\n'
    '答：这次的目标是先在按标本验证的规则下，给出可信的基线成绩。结果显示瓶颈主要在每种矿物的标本数量，换更复杂的模型不一定能解决。\n'
    '问：被问到不会的问题时\n'
    '答：这个问题我们这次没有研究，会后我们查一下再回复您。')
text(s, 0.67, 1.6, 11.9, 3.8, [
    ('约七成', 23.25, BLUE, True),
    ('类别扩大到 1,669 类后，按标本留出的最好正确率约 71%', 20.25, NAVY, False), ('', 10, NAVY, False),
    ('简单方法够用', 23.25, BLUE, True),
    ('余弦近邻与 SVM 相当；错误多在同族、固溶体矿物之间', 20.25, NAVY, False), ('', 10, NAVY, False),
    ('数据是瓶颈', 23.25, BLUE, True),
    ('每种矿物的独立标本越多越准，比换模型更有效', 20.25, NAVY, False)], space=10)
rect(s, 0.71, 5.65, 11.9, 0.95, PALE)
text(s, 1.0, 5.65, 11.4, 0.95, [('下一步：为同一矿物收集更多独立标本；用其他实验室、其他仪器的光谱做外部检验', 18, SLATE, False)],
     anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print('wrote', OUT, len(prs.slides), 'slides')
